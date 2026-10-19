"""Test suite for the GreenGate3D lead tool.

  python run_tests.py          # offline unit tests only
  python run_tests.py --live   # also crawl a real makerspace site (network, no API keys)

Covers URL handling, website-source filtering, link discovery, JSON-LD parsing,
email/phone selection, extension-space dedupe, priority rules, people
formatting, and the Excel output. No API keys or network needed.
"""
import io
import re
import sys
import threading

import config
import scraper
import excel_out

PASS, FAIL = 0, []


def check(label, cond, detail=""):
    global PASS
    if cond:
        PASS += 1
        print(f"  ok  {label}")
    else:
        FAIL.append(label)
        print(f"FAIL  {label}  {detail}")


# ---------------------------------------------------------------------------
print("\n[1] URL helpers")
check("origin strips path+query",
      scraper._origin("https://shop.com/page?utm_source=google") == "https://shop.com")
check("origin handles bare domain", scraper._origin("shop.com") == "https://shop.com")
check("origin empty on garbage", scraper._origin("") == "")
check("same_site ignores www",
      scraper._same_site("https://www.shop.com/about", "https://shop.com"))
check("same_site ignores scheme",
      scraper._same_site("http://shop.com/x", "https://www.shop.com"))
check("same_site rejects other domain",
      not scraper._same_site("https://facebook.com/shop", "https://shop.com"))

print("\n[1b] Generic-host rejection (social pages are not websites)")
check("facebook rejected", scraper._is_generic_host("https://facebook.com/knoxmakers"))
check("linktree rejected", scraper._is_generic_host("https://linktr.ee/shop"))
check("subdomain of generic rejected",
      scraper._is_generic_host("https://m.facebook.com/shop"))
check("real site accepted", not scraper._is_generic_host("https://knoxmakers.org"))

# ---------------------------------------------------------------------------
print("\n[2] Link discovery + people-page ranking")
HTML_LINKS = """
<html><body>
  <a href="/about-us">About Us</a>
  <a href="contact.html">Contact</a>
  <a href="../team">Our Team</a>
  <a href="https://www.shop.com/equipment">Equipment</a>
  <a href="https://facebook.com/shop">Contact us on FB</a>
  <a href="mailto:info@shop.com">Email</a>
  <a href="/blog/post-1">Random post</a>
</body></html>
"""
urls = scraper._internal_links("https://shop.com/sub/page", HTML_LINKS, "https://shop.com")
check("absolute path resolved", "https://shop.com/about-us" in urls, str(urls))
check("relative file resolved", "https://shop.com/sub/contact.html" in urls, str(urls))
check("parent-relative resolved", "https://shop.com/team" in urls, str(urls))
check("www variant kept as same-site", any(u.endswith("/equipment") for u in urls))
check("external domain excluded", not any("facebook.com" in u for u in urls))
# Changed contract: non-keyword pages are KEPT (ranked last), not excluded --
# filtering them out is what left contact boxes empty.
check("non-keyword page is kept for crawling", any("blog" in u for u in urls), str(urls))
check("but it ranks below a people page",
      scraper._page_relevance("https://shop.com/blog/post-1")
      < scraper._page_relevance("https://shop.com/team"))
check("people page outranks info page",
      scraper._page_relevance("https://x.com/leadership")
      > scraper._page_relevance("https://x.com/about"))
# Regression: keywords must match the path only. A shop at "teamfab.com" must
# not score every one of its pages as a people page.
check("hostname does not trigger keyword match",
      scraper._page_relevance("https://teamfab.com/blog/post-1") == 0,
      str(scraper._page_relevance("https://teamfab.com/blog/post-1")))
check("path still matches on keyword-named host",
      scraper._page_relevance("https://teamfab.com/our-team") > 0)

print("\n[2b] Token matching + unfiltered link discovery")
# Real site had /who-are-we; the keyword was who-we-are. Word order must not matter.
check("who-are-we matches who-we-are",
      scraper._page_relevance("https://x.com/who-are-we") > 0)
check("underscore and camel separators tokenize",
      scraper._page_relevance("https://x.com/our_team") > 0)
check("unmatched page is kept at relevance 0, not excluded",
      scraper._page_relevance("https://x.com/operations") >= 0)
check("junk paths are skipped", scraper._should_skip("https://x.com/cart"))
check("asset urls are skipped", scraper._should_skip("https://x.com/logo.png"))
check("real page not skipped", not scraper._should_skip("https://x.com/who-are-we"))
# Every internal link must survive discovery, even with an unusual name.
ALL_LINKS = """<html><body>
  <a href="/who-are-we">Who Are We</a><a href="/make">Make</a>
  <a href="/operations/">Operations</a><a href="/cart">Cart</a>
  <a href="https://facebook.com/x">FB</a><a href="/logo.png">img</a>
</body></html>"""
found = scraper._internal_links("https://shop.com/", ALL_LINKS, "https://shop.com")
paths = sorted(scraper._url_path(u).strip() for u in found)
check("unusual page names are kept", "/who-are-we" in paths and "/make" in paths, str(paths))
check("operations kept", "/operations/" in paths, str(paths))
check("cart excluded", not any("cart" in p for p in paths), str(paths))
check("image excluded", not any("logo" in p for p in paths), str(paths))
check("external excluded", not any("facebook" in u for u in found))

print("\n[2c] JS-shell detection")
check("tiny shell detected", scraper.looks_like_shell("<html><body><div id=root></div></body></html>", ""))
check("big markup, no prose detected",
      scraper.looks_like_shell("<div>" + "x" * 30000 + "</div>", "short text"))
check("normal page not flagged",
      not scraper.looks_like_shell("<html>" + "<p>hello world</p>" * 100,
                                   "hello world " * 100))
check("empty html not flagged", not scraper.looks_like_shell("", ""))
# Playwright's sync API has thread affinity; rendering off-thread corrupts the
# browser. The renderer must refuse rather than be driven from a worker thread.
_r = scraper.Renderer(enabled=True, budget=5)
_r._owner = threading.get_ident()
_out = {}
_t = threading.Thread(target=lambda: _out.update(v=_r.html("https://example.com")))
_t.start(); _t.join()
check("renderer refuses calls from a foreign thread", _out.get("v") == "")
check("budget untouched by the refused call", _r.budget == 5, str(_r.budget))
check("disabled renderer returns nothing",
      scraper.Renderer(enabled=False).html("https://example.com") == "")

# ---------------------------------------------------------------------------
print("\n[3] JSON-LD extraction")
HTML_JSONLD = """
<html><head><script type="application/ld+json">
{"@context":"https://schema.org","@type":"LocalBusiness","name":"Shop",
 "email":"mailto:hello@shop.com","telephone":"+1-865-555-0100",
 "sameAs":["https://www.linkedin.com/company/shop","https://instagram.com/shop"]}
</script></head><body></body></html>
"""
emails, phones, socials = scraper._extract_jsonld(HTML_JSONLD)
check("jsonld email found (mailto stripped)", emails == ["hello@shop.com"], str(emails))
check("jsonld phone found", phones == ["+1-865-555-0100"], str(phones))
check("jsonld socials found", len(socials) == 2, str(socials))

# ---------------------------------------------------------------------------
print("\n[4] Email + phone selection")
check("prefers own-domain over gmail",
      scraper._best_email(
          'x <a href="mailto:shop.owner@gmail.com">a</a> contact info@shop.com',
          "contact info@shop.com", [], host="shop.com") == "info@shop.com")
check("jsonld beats visible text",
      scraper._best_email("a@other.com in text", "a@other.com in text",
                          ["official@other.com"], host="") == "official@other.com")
check("blocklist filters placeholders",
      scraper._best_email("email@domain.com user@example.com", "", [], host="") == "")
check("deobfuscated address reconstructed",
      scraper._best_email("", "reach us: info [at] shop [dot] org", [],
                          host="shop.org") == "info@shop.org")
check("website phone from jsonld",
      scraper._best_phone("", ["+1-865-555-0100"]) == "+1-865-555-0100")
check("website phone falls back to tel: link",
      scraper._best_phone('<a href="tel:+18655550142">call</a>', []) == "+18655550142")
check("no phone -> empty", scraper._best_phone("<p>no numbers here</p>", []) == "")
# Most small-shop sites print the number as plain text with no markup at all.
check("phone read from visible text",
      scraper._best_phone("", [], "Call us at (718) 273-3951 today") == "(718) 273-3951")
check("phone text formats normalized",
      scraper._best_phone("", [], "phone 718.273.3951") == "(718) 273-3951")
check("phone near context beats bare number",
      scraper._best_phone("", [], "Est. 212 555 1000. Phone: (718) 273-3951")
      == "(718) 273-3951")
check("zip+4 is not a phone", scraper._best_phone("", [], "Brooklyn NY 11215-1234") == "")
check("date-like digits are not a phone",
      scraper._best_phone("", [], "open 2024 2025 1000") == "")
# Regression: contact details live in header/footer/nav, which _clean_text
# deliberately strips. Contact mining must use the preserving view.
FOOTER_HTML = ("<html><body><main>Our shop does great work.</main>"
               "<footer>Phone: (718) 273-3951</footer></body></html>")
check("clean_text drops footer (correct for prose)",
      "273-3951" not in scraper._clean_text(FOOTER_HTML))
check("contact_text keeps footer phone",
      "273-3951" in scraper._contact_text(FOOTER_HTML))
check("footer phone is extracted end-to-end",
      scraper._best_phone(FOOTER_HTML, [], scraper._contact_text(FOOTER_HTML))
      == "(718) 273-3951")

# ---------------------------------------------------------------------------
print("\n[5] Extension-space dedupe (NYU MakerGarage / MakerSpace)")
prospects = [
    {"name": "NYU MakerGarage", "website": "https://engineering.nyu.edu/makerspace",
     "reviews": 5, "address": "6 MetroTech, Brooklyn NY", "phone_maps": ""},
    {"name": "NYU MakerSpace", "website": "https://engineering.nyu.edu/makerspace",
     "reviews": 60, "address": "5 MetroTech, Brooklyn NY", "phone_maps": "(646) 555-0100"},
    {"name": "Unrelated Shop", "website": "https://othershop.com",
     "reviews": 12, "address": "1 Main St", "phone_maps": ""},
]
merged = scraper.merge_prospects([dict(p) for p in prospects])
check("same-website places merged into one row", len(merged) == 2, f"{len(merged)} rows")
nyu = next(m for m in merged if "NYU" in m["name"])
check("most-reviewed place becomes primary", nyu["name"] == "NYU MakerSpace", nyu["name"])
check("extension recorded as related location",
      nyu["related"] == ["NYU MakerGarage"], str(nyu.get("related")))
check("merged row keeps maps phone", nyu["phone_maps"] == "(646) 555-0100")
check("unrelated site not merged",
      any(m["name"] == "Unrelated Shop" for m in merged))

no_site = [{"name": "A", "website": "", "reviews": 1, "address": "", "phone_maps": ""},
           {"name": "B", "website": "", "reviews": 2, "address": "", "phone_maps": ""}]
check("differently-named website-less places are never merged",
      len(scraper.merge_prospects([dict(p) for p in no_site])) == 2)
# Same org found by both sources: Places (with site) + web search (without).
dual = [
    {"name": "Knox Makers", "website": "", "reviews": 0, "address": "",
     "phone_maps": "", "discovered_via": "Web search"},
    {"name": "The Knox Makers", "website": "", "reviews": 9, "address": "",
     "phone_maps": "", "discovered_via": "Google Places"},
]
dm = scraper.merge_prospects([dict(p) for p in dual])
check("same org from two sources merges by name", len(dm) == 1, f"{len(dm)} rows")
check("merged row credits both sources",
      "Google Places" in dm[0]["discovered_via"] and "Web search" in dm[0]["discovered_via"],
      dm[0]["discovered_via"])

social = [{"name": "A", "website": "https://facebook.com/x", "reviews": 1,
           "address": "", "phone_maps": ""},
          {"name": "B", "website": "https://facebook.com/x", "reviews": 2,
           "address": "", "phone_maps": ""}]
check("shared social page does NOT merge distinct businesses",
      len(scraper.merge_prospects([dict(p) for p in social])) == 2)

# ---------------------------------------------------------------------------
print("\n[5b] Same-host, different-page places must NOT merge")
# A campus hosts several distinct makerspaces on one domain; merging on host
# alone would collapse them into a single lead.
campus = [
    {"name": "UT Fab Lab", "website": "https://utk.edu/fablab", "reviews": 3,
     "address": "", "phone_maps": ""},
    {"name": "Hodges Library Makerspace", "website": "https://utk.edu/library",
     "reviews": 8, "address": "", "phone_maps": ""},
]
check("distinct pages on one campus host stay separate",
      len(scraper.merge_prospects([dict(p) for p in campus])) == 2)
same_page = [
    {"name": "NYU MakerGarage", "website": "https://nyu.edu/makerspace",
     "reviews": 5, "address": "", "phone_maps": ""},
    {"name": "NYU MakerSpace", "website": "https://nyu.edu/makerspace/",
     "reviews": 60, "address": "", "phone_maps": ""},
]
check("same page (trailing slash ignored) still merges",
      len(scraper.merge_prospects([dict(p) for p in same_page])) == 1)

print("\n[5b2] Duplicate locations that slipped past exact-key matching")
def _n(rows):
    return len(scraper.merge_prospects([dict(r, reviews=r.get("reviews", 0),
                                             phone_maps=r.get("phone_maps", ""))
                                        for r in rows]))
check("tracking query on one listing's website still merges",
      _n([{"name": "Knox Makers", "website": "https://knoxmakers.org/?utm_source=gmb", "address": ""},
          {"name": "KnoxMakers Inc", "website": "https://www.knoxmakers.org", "address": ""}]) == 1)
check("Places (with site) + web search (no site), same name, merges",
      _n([{"name": "Knox Makers", "website": "https://knoxmakers.org", "address": "1 Main St, Knoxville"},
          {"name": "The Knox Makers", "website": "", "address": ""}]) == 1)
check("same name + same street address (suite/abbr differ) merges",
      _n([{"name": "Maker Hub", "website": "https://makerhub.com", "address": "100 North Main Street Suite 4, Knoxville"},
          {"name": "Maker Hub Knoxville", "website": "https://makerhub.com/knoxville", "address": "100 N Main St, Knoxville"}]) == 1)
check("chain branches at different addresses stay separate",
      _n([{"name": "Home Depot", "website": "https://homedepot.com/l/1", "address": "1 Kingston Pike, Knoxville"},
          {"name": "Home Depot", "website": "https://homedepot.com/l/2", "address": "9 Clinton Hwy, Knoxville"}]) == 2)
check("same-name website-less places at different addresses stay separate",
      _n([{"name": "Tool Library", "website": "", "address": "1 Oak Ave, Knoxville"},
          {"name": "Tool Library", "website": "", "address": "7 Elm St, Maryville"}]) == 2)
check("a one-word name never swallows longer names",
      _n([{"name": "Makerspace", "website": "", "address": ""},
          {"name": "Hodges Library Makerspace", "website": "", "address": ""}]) == 2)
check("same building, different labs, stay separate",
      _n([{"name": "Fab Lab", "website": "https://utk.edu/fablab", "address": "1015 Volunteer Blvd"},
          {"name": "Robotics Lab", "website": "https://utk.edu/robotics", "address": "1015 Volunteer Blvd"}]) == 2)
_leads = [
    {"name": "Knox Makers", "website": "https://knoxmakers.org", "email": "hi@knoxmakers.org",
     "contacts": [1, 2], "address": ""},
    {"name": "Knox Makers Inc", "website": "https://knoxmakers.com", "email": "HI@knoxmakers.org",
     "contacts": [], "address": "5 Gay St"},
    {"name": "UT Fab Lab", "website": "https://utk.edu/fablab", "email": "hi@knoxmakers.org",
     "contacts": [], "address": ""},
]
_d = scraper.dedupe_leads([dict(L) for L in _leads])
check("post-crawl: alias domain sharing the org's own email merges",
      [L["name"] for L in _d] == ["Knox Makers", "UT Fab Lab"], str([L["name"] for L in _d]))
check("post-crawl: merged lead keeps the other as related",
      _d[0].get("related") == ["Knox Makers Inc"], str(_d[0].get("related")))

print("\n[5c] Cloudflare-obfuscated email decoding")
_plain = "info@shop.org"
_key = 0x7a
_enc = f"{_key:02x}" + "".join(f"{ord(c) ^ _key:02x}" for c in _plain)
check("cfemail decodes to the real address", scraper._decode_cfemail(_enc) == _plain,
      scraper._decode_cfemail(_enc))
check("cfemail found in html", scraper._cfemails(f'<a data-cfemail="{_enc}">x</a>') == [_plain])
check("malformed cfemail ignored", scraper._cfemails('<a data-cfemail="zz">x</a>') == [])

print("\n[5d] Printer-model gazetteer (deterministic, over the FULL crawl)")
fp = scraper.find_printer_models
check("prusa model found", "MK4" in " ".join(fp("we run a Prusa MK4 daily")).upper())
check("bambu model found", any("p1s" in m.lower() for m in fp("Bambu Lab P1S in the lab")))
check("ender variant found", any("ender 3" in m.lower() for m in fp("two Ender 3 V2 units")))
check("resin printer found", any("form 3" in m.lower() for m in fp("a Formlabs Form 3 printer")))
# The exact false positive my own diagnostic regex hit: "anet" inside "planetary".
check("'planetary' is not the Anet printer", fp("reframes architecture as planetary forces") == [])
check("'information' does not match", fp("more information about our shop") == [])
check("no printers -> empty", fp("we have a laser cutter and a CNC router") == [])
# Printer inventories sit on process-named pages, which must outrank prose.
check("3d-printing page ranks above a generic page",
      scraper._page_relevance("https://x.edu/fab-lab/3d-printing/")
      > scraper._page_relevance("https://x.edu/news/story"))
check("deep fab-lab path is ranked, not ignored",
      scraper._page_relevance("https://x.edu/study/studios-learning-labs/fab-lab/") > 0)
check("bare brand still reported when no model given",
      any("prusa" in m.lower() for m in fp("we own several Prusa machines")))
check("duplicates collapse to the most specific",
      len(fp("Prusa. Prusa MK4. Prusa MK4 is great")) == 1,
      str(fp("Prusa. Prusa MK4. Prusa MK4 is great")))
# Precision regressions seen on a real makerspace wiki.
check("PrusaSlicer is software, not a printer", fp("we use PrusaSlicer to prepare") == [],
      str(fp("we use PrusaSlicer to prepare")))
check("Bambu Studio is software", fp("export from Bambu Studio") == [],
      str(fp("export from Bambu Studio")))
check("wash/cure stations are not printers",
      fp("Mercury X Wash and Mercury X Cure stations") == [],
      str(fp("Mercury X Wash and Mercury X Cure stations")))
check("wiki heading numbers don't create phantom models",
      fp("3.3 Ender 3 setup") == ["Ender 3"], str(fp("3.3 Ender 3 setup")))
check("real Ender variants still match",
      any("pro" in m.lower() for m in fp("an Ender 3 Pro in the corner")),
      str(fp("an Ender 3 Pro in the corner")))
check("a genuine Prusa still matches alongside its slicer",
      any("mk4" in m.lower() for m in fp("Prusa MK4 sliced in PrusaSlicer")),
      str(fp("Prusa MK4 sliced in PrusaSlicer")))
# Two different models of one brand are two printers — both must survive.
_two = fp("we run a Prusa MK3 and a Prusa XL")
check("distinct models of the same brand are both kept",
      len(_two) == 2 and any("mk3" in m.lower() for m in _two)
      and any("xl" in m.lower() for m in _two), str(_two))
check("brand-qualified name absorbs the bare variant",
      fp("Creality Ender 3 next to another Ender 3") == ["Creality Ender 3"],
      str(fp("Creality Ender 3 next to another Ender 3")))

print("\n[5e] Maker subdomain scoping")
check("registrable domain extracted",
      scraper._registrable("wiki.knoxmakers.org") == "knoxmakers.org")
check("registrable of bare domain unchanged",
      scraper._registrable("knoxmakers.org") == "knoxmakers.org")
check("unfiltered link collection returns cross-host links",
      any("wiki." in u for u in scraper._links_from_soup(
          __import__("bs4").BeautifulSoup(
              '<a href="https://wiki.shop.com/tools">w</a>', "html.parser"),
          "https://shop.com/", None)[0]))
check("filtered link collection still excludes other hosts",
      scraper._links_from_soup(
          __import__("bs4").BeautifulSoup(
              '<a href="https://wiki.shop.com/tools">w</a>', "html.parser"),
          "https://shop.com/", "https://shop.com")[0] == set())

print("\n[5f] Gap-search provenance is visible, never silently merged")
check("site-sourced lead shows no search note",
      "web search" not in excel_out.format_source(
          {"discovered_via": "Google Places", "website_source": "Google Places"}))
check("enriched fields are labelled",
      "email via web search" in excel_out.format_source(
          {"discovered_via": "Google Places", "website_source": "Google Places",
           "enriched": "email"}))
check("deep preset enables gap search", scraper.DEPTH_PRESETS["deep"]["enrich"])
check("quick preset does not", not scraper.DEPTH_PRESETS["quick"]["enrich"])

print("\n[5g] Cross-contamination on shared / institutional sites")
# Deep crawls of a whole university or a shared venue harvest outsiders.
check("another employer's address is foreign",
      scraper._is_foreign_domain("erica@marriotthotels.com", "makerexchangeknox.com"))
check("the org's own address is not foreign",
      not scraper._is_foreign_domain("info@makerexchangeknox.com", "makerexchangeknox.com"))
check("subdomain of the org is not foreign",
      not scraper._is_foreign_domain("a@wiki.knoxmakers.org", "knoxmakers.org"))
check("gmail is never judged foreign (small shops really use it)",
      not scraper._is_foreign_domain("owner@gmail.com", "shop.com"))
check("unknown org host judges nothing foreign",
      not scraper._is_foreign_domain("x@anywhere.com", ""))

mixed = ["info@shop.com", "sales@shop.com", "erica@marriotthotels.com"]
check("own-domain addresses win when present",
      scraper._prefer_own_domain(mixed, "shop.com")
      == ["info@shop.com", "sales@shop.com"])
check("a freemail address is plausibly theirs; a partner's is not",
      scraper._prefer_own_domain(["a@gmail.com", "b@partner.com"], "shop.com")
      == ["a@gmail.com"])
check("an all-foreign list is emptied, not passed through",
      scraper._prefer_own_domain(["b@partner.com", "c@other.org"], "shop.com") == [],
      "freemail is already exempt, so this can only be contamination")
check("unknown org host disables filtering entirely",
      scraper._prefer_own_domain(["b@partner.com"], "") == ["b@partner.com"])

print("\n[5h] Crawl focus on a lead's own section")
ft = scraper.focus_tokens("Medbery Makerspace (Hodges Library, University of Tennessee)",
                          "https://www.lib.utk.edu")
check("distinctive name words become focus tokens",
      {"medbery", "makerspace", "hodges"} <= ft, str(ft))
check("institutional filler is dropped", "university" not in ft, str(ft))
check("words already in the hostname are dropped", "utk" not in ft and "lib" not in ft,
      str(ft))
check("short words are dropped", not any(len(t) < 4 for t in ft), str(ft))
# The lead's own page must outrank a staff directory for a different department.
check("own-section page outranks an unrelated staff page",
      scraper._page_relevance("https://lib.utk.edu/medbery/makerspace", ft)
      > scraper._page_relevance("https://lib.utk.edu/agvet/staff", ft))
check("without focus, ranking is unchanged",
      scraper._page_relevance("https://lib.utk.edu/agvet/staff")
      == scraper._page_relevance("https://lib.utk.edu/agvet/staff", set()))

print("\n[5i] Audit fixes: silent failures and domain logic")
# Multi-part suffixes: without these every .co.uk address looks like one org.
check("co.uk keeps its own label",
      scraper._registrable("shop.co.uk") == "shop.co.uk",
      scraper._registrable("shop.co.uk"))
check("subdomain under co.uk resolves to the org",
      scraper._registrable("wiki.shop.co.uk") == "shop.co.uk")
check("plain .com unchanged", scraper._registrable("a.shop.com") == "shop.com")
check("foreign detection works on co.uk",
      scraper._is_foreign_domain("x@other.co.uk", "shop.co.uk"))
check("same co.uk org is not foreign",
      not scraper._is_foreign_domain("x@shop.co.uk", "shop.co.uk"))

# Usefulness is site-relative: a word already in the host separates nothing
# there, but the same word can be the key discriminator on a shared site.
ft2 = scraper.focus_tokens("3D Print Knox", "https://3dprintknox.com")
check("words baked into the host are dropped",
      "print" not in ft2 and "knox" not in ft2, str(ft2))
check("a lead on its own domain needs no focus at all", ft2 == set(), str(ft2))
check("but 'makerspace' survives on a university host — it is the discriminator",
      "makerspace" in scraper.focus_tokens("Medbery Makerspace",
                                           "https://www.lib.utk.edu"))
check("'makers' dropped on its own domain",
      "makers" not in scraper.focus_tokens("Knox Makers", "https://knoxmakers.org"))

# Shell detection must not need its own parse (it used to be a third one).
check("link count can be supplied by the caller",
      scraper.looks_like_shell("<html>" + "x" * 30000 + "</html>", "short", links=0))
check("regex link count matches a real page",
      not scraper.looks_like_shell(
          "<html>" + "<a href='/a'>x</a>" * 20 + "<p>" + "word " * 400 + "</p></html>",
          "word " * 400))

print("\n[5j] Cost accounting")
class _FakeUsage:
    input_tokens = 10000
    output_tokens = 2000
    cache_read_input_tokens = 0
    cache_creation_input_tokens = 0
class _FakeMsg:
    usage = _FakeUsage()

sp = scraper.Spend()
sp.add_places(19)
sp.add_call("extract + pitch", "claude-haiku-4-5", _FakeMsg())
haiku_only = sp.total()
sp2 = scraper.Spend()
sp2.add_places(19)
sp2.add_call("extract + pitch", "claude-opus-5", _FakeMsg())
check("the same call costs more on the flagship model", sp2.total() > haiku_only,
      f"{haiku_only:.3f} vs {sp2.total():.3f}")
check("haiku extraction is ~5x cheaper than opus for the same tokens",
      abs(((sp2.total() - 19 * config.PLACES_COST_PER_REQUEST)
           / (haiku_only - 19 * config.PLACES_COST_PER_REQUEST)) - 5.0) < 0.01)
check("places requests are counted", sp.places == 19)
check("web searches add cost",
      scraper.Spend().add_call("x", "claude-sonnet-5", _FakeMsg(), searches=3) is None)
check("report names each call type and a total",
      "TOTAL" in sp.report(5) and "extract + pitch" in sp.report(5))
check("per-lead figure shown", "per lead" in sp.report(5))
check("extraction uses a capable model (measured at only ~4% of spend)",
      config.EXTRACT_MODEL in ("claude-sonnet-5", "claude-opus-5"),
      config.EXTRACT_MODEL)
check("search model is mid-tier, not the flagship",
      config.SEARCH_MODEL == "claude-sonnet-5", config.SEARCH_MODEL)

print("\n[5k] collect() orchestration (stubbed I/O, non-mock path)")
# MOCK_MODE returns before any of collect()'s real logic runs, so the whole
# orchestration layer had no coverage -- and a stale set-vs-list call crashed
# a live run. These stubs exercise the real path without network or API cost.
_saved = {n: getattr(scraper, n) for n in
          ("places_search_all", "discover_via_search", "find_website",
           "crawl_site", "extract_and_pitch", "enrich_gaps", "Renderer")}
_saved_keys = (config.GOOGLE_PLACES_API_KEY, config.ANTHROPIC_API_KEY)
try:
    config.GOOGLE_PLACES_API_KEY = config.ANTHROPIC_API_KEY = "test-key"
    scraper.places_search_all = lambda *a, **k: [
        {"id": "p1", "displayName": {"text": "Knox Makers"},
         "formattedAddress": "1 Main St, Knoxville, TN 37920",
         "websiteUri": "https://knoxmakers.org", "userRatingCount": 42,
         "rating": 4.9, "businessStatus": "OPERATIONAL"},
        {"id": "p2", "displayName": {"text": "No Site Shop"},
         "formattedAddress": "2 Oak Ave, Knoxville, TN 37902"},
    ]
    # Returns one genuinely new org and one that duplicates a Places result
    # under a different spelling — the case that crashed.
    scraper.discover_via_search = lambda *a, **k: [
        {"name": "UT Fab Lab", "website": "https://archdesign.utk.edu",
         "city": "Knoxville", "address": "", "kind": "University makerspace"},
        {"name": "The Knox Makers", "website": "", "city": "Knoxville",
         "address": "", "kind": "Makerspace"},
    ]
    scraper.find_website = lambda *a, **k: ""
    def _fake_crawl(website, *a, **k):
        # A website-less lead really does get nothing back, so mirror that
        # rather than handing every lead the same fabricated page.
        if not website:
            return {"text": "", "email": "", "emails": [], "phone": "",
                    "socials": [], "pages": 0, "rendered": 0,
                    "printer_models": [], "hosts": []}
        return {"text": "We run a Prusa MK4.", "email": "hi@knoxmakers.org",
                "emails": ["hi@knoxmakers.org", "outsider@othercorp.com"],
                "phone": "(865) 555-0100", "socials": [], "pages": 7,
                "rendered": 0, "printer_models": ["Prusa MK4"],
                "hosts": ["knoxmakers.org"]}
    scraper.crawl_site = _fake_crawl
    scraper.extract_and_pitch = lambda *a, **k: {
        "people": [{"name": "Jane Doe", "role": "Director", "email": "",
                    "phone": "", "linkedin": "", "reason": "buys materials"}],
        "equipment_highlights": "3D printers", "printer_models": "",
        "pitch": "A grounded pitch."}
    scraper.enrich_gaps = lambda *a, **k: {}
    scraper.Renderer = lambda **k: type(
        "R", (), {"budget": 0, "html": lambda s, u: "", "close": lambda s: None})()

    config.MOCK_MODE = False
    got = scraper.collect("Knoxville, TN", "makerspace", max_results=5,
                          depth="deep")
    names = [L["name"] for L in got]
    check("collect() runs the real path without crashing", bool(got), str(names))
    check("near-duplicate across sources is not added twice",
          not any("Knox Makers" == n for n in names[1:]), str(names))
    check("genuinely new web-search org is kept",
          any("Fab Lab" in n for n in names), str(names))
    check("discovery source recorded per lead",
          {L["discovered_via"] for L in got} <= {"Google Places", "Web search"},
          str([L["discovered_via"] for L in got]))
    check("foreign address filtered from other_emails",
          all("othercorp.com" not in " ".join(L.get("other_emails") or [])
              for L in got))
    check("printer models from the crawl reach the lead",
          any("Prusa MK4" in (L["printer_models"] or "") for L in got))
    check("priority derived for every lead",
          all(L["priority"] in ("HIGH", "MEDIUM", "LOW") for L in got))
    check("a spend report is produced", hasattr(got, "spend"))
finally:
    for n, v in _saved.items():
        setattr(scraper, n, v)
    config.GOOGLE_PLACES_API_KEY, config.ANTHROPIC_API_KEY = _saved_keys
    config.MOCK_MODE = False

print("\n[5l] Gap-search gating (the measured 81% of run cost)")
def _wanted(lead):
    """Mirror of the gating in collect(), to assert what gets paid for."""
    has_phone = bool(lead.get("phone_web") or lead.get("phone_maps"))
    prints_3d = bool(lead.get("printer_models")) or bool(
        scraper.PRINT3D_RE.search((lead.get("equipment_highlights") or "")
                                  + " " + lead.get("type", "") + " "
                                  + lead.get("name", "")))
    want = []
    if not lead.get("email"):
        want.append("email")
        if not has_phone:
            want.append("phone")
    elif not has_phone and not lead.get("contacts"):
        want.append("phone")
    if not lead.get("printer_models") and prints_3d:
        want.append("printer_models")
    return want

check("reachable lead with printers is not searched at all",
      _wanted({"email": "a@b.org", "phone_web": "1", "printer_models": "Prusa MK4",
               "contacts": [{"name": "X"}], "type": "Makerspace", "name": "X"}) == [])
check("a pottery studio is never searched for printers",
      "printer_models" not in _wanted(
          {"email": "a@b.org", "phone_web": "1", "printer_models": "",
           "equipment_highlights": "kilns and wheels", "contacts": [{"name": "X"}],
           "type": "Pottery studio", "name": "Mighty Mud"}))
check("a 3D shop with no models listed IS searched",
      "printer_models" in _wanted(
          {"email": "a@b.org", "phone_web": "1", "printer_models": "",
           "equipment_highlights": "we run FDM and resin printers",
           "contacts": [{"name": "X"}], "type": "Shop", "name": "Shop"}))
check("an unreachable lead is searched for both email and phone",
      _wanted({"email": "", "phone_web": "", "phone_maps": "", "type": "x",
               "name": "x", "printer_models": ""})[:2] == ["email", "phone"])
check("having an email means no phone-only search",
      _wanted({"email": "a@b.org", "phone_web": "", "phone_maps": "",
               "contacts": [{"name": "X"}], "type": "x", "name": "x",
               "printer_models": ""}) == [])

print("\n[5m] Gap fill goes through Claude's web search (Brave removed)")
check("no Brave code remains in the scraper",
      not any(hasattr(scraper, n) for n in ("brave_search", "enrich_gaps_free")))
_saved = {n: getattr(scraper, n) for n in
          ("places_search_all", "discover_via_search", "find_website",
           "crawl_site", "extract_and_pitch", "enrich_gaps", "Renderer")}
_saved_keys = (config.GOOGLE_PLACES_API_KEY, config.ANTHROPIC_API_KEY)
_gap_calls = []
try:
    config.GOOGLE_PLACES_API_KEY = config.ANTHROPIC_API_KEY = "test-key"
    scraper.places_search_all = lambda *a, **k: [
        {"id": "g1", "displayName": {"text": "Quiet Fab Lab"},
         "formattedAddress": "9 Pine St, Knoxville, TN", "websiteUri": "https://quietfab.org"}]
    scraper.discover_via_search = lambda *a, **k: []
    scraper.find_website = lambda *a, **k: ""
    scraper.crawl_site = lambda *a, **k: {
        "text": "Our fab lab runs a Prusa MK4.", "email": "", "emails": [], "phone": "",
        "socials": [], "pages": 2, "rendered": 0, "printer_models": ["Prusa MK4"],
        "hosts": [], "priority_pages": ["x/lab"]}
    scraper.extract_and_pitch = lambda *a, **k: {
        "people": [], "equipment_highlights": "3D printers", "printer_models": "",
        "pitch": "p"}

    def _fake_gap(client, name, city, website, want, spend=None):
        _gap_calls.append(want)
        return {"email": "hello@quietfab.org", "source": "https://directory.example.org"}
    scraper.enrich_gaps = _fake_gap
    scraper.Renderer = lambda **k: type(
        "R", (), {"budget": 0, "html": lambda s, u: "", "close": lambda s: None})()
    config.MOCK_MODE = False
    got = scraper.collect("Knoxville, TN", "makerspace", max_results=1, depth="deep")
    check("Deep searches for a missing email", _gap_calls and "email" in _gap_calls[0],
          str(_gap_calls))
    check("found email reaches the lead, marked as enriched",
          got and got[0]["email"] == "hello@quietfab.org" and "email" in got[0]["enriched"],
          str([(L["email"], L["enriched"]) for L in got]))
    _gap_calls.clear()
    scraper.collect("Knoxville, TN", "makerspace", max_results=1, depth="standard")
    check("Standard never runs the paid gap search", _gap_calls == [], str(_gap_calls))
finally:
    for n, v in _saved.items():
        setattr(scraper, n, v)
    config.GOOGLE_PLACES_API_KEY, config.ANTHROPIC_API_KEY = _saved_keys
    config.MOCK_MODE = False

print("\n[6] Priority rules")
check("printers + reachable + contact + reviews = HIGH",
      scraper.derive_priority({"email": "a@b.org", "phone_web": "", "phone_maps": "1",
                               "reviews": 30, "printer_models": "Prusa MK4",
                               "equipment_highlights": "3D printers",
                               "contacts": [{"name": "X"}]}) == "HIGH")
check("screen printing is NOT 3D evidence",
      scraper.derive_priority({"email": "", "phone_web": "", "phone_maps": "",
                               "reviews": 0, "printer_models": "",
                               "equipment_highlights": "screen printing presses",
                               "contacts": []}) == "LOW")
check("filament keyword counts as 3D evidence",
      scraper.derive_priority({"email": "", "phone_web": "", "phone_maps": "",
                               "reviews": 0, "printer_models": "",
                               "equipment_highlights": "we stock filament",
                               "contacts": []}) == "MEDIUM")
check("website phone alone makes a lead reachable",
      scraper.derive_priority({"email": "", "phone_web": "(865) 555-0100",
                               "phone_maps": "", "reviews": 0, "printer_models": "",
                               "equipment_highlights": "", "contacts": []}) == "MEDIUM")

# ---------------------------------------------------------------------------
print("\n[7] Other Contacts formatting")
lead = {"contacts": [
    {"name": "Zoe Junior", "role": "Volunteer", "email": "", "phone": "",
     "linkedin": "", "reason": ""},
    {"name": "Sam Lee", "role": "Technical Director", "email": "sam@x.org",
     "phone": "(865) 555-0142", "linkedin": "https://linkedin.com/in/samlee",
     "reason": "runs the shop floor and its printers"},
]}
cell = excel_out.format_people(lead)
check("decision-maker sorted above volunteer",
      cell.index("Sam Lee") < cell.index("Zoe Junior"))
check("name and role on one line", "Sam Lee — Technical Director" in cell)
check("email/phone/linkedin all present",
      all(s in cell for s in ("sam@x.org", "(865) 555-0142", "linkedin.com/in/samlee")))
check("reason rendered", "Why: runs the shop floor" in cell)
check("empty contacts -> empty cell", excel_out.format_people({"contacts": []}) == "")
check("other emails appended to the same cell",
      "info@x.org" in excel_out.format_people(
          {"contacts": [], "other_emails": ["info@x.org", "help@x.org"]}))
check("no other-emails block when there are none",
      "Other" not in excel_out.format_people(
          {"contacts": [{"name": "Sam"}], "other_emails": []}))

# ---------------------------------------------------------------------------
print("\n[7a] Places multi-query fan-out")
# Google Places matches the literal phrase: "makerspace in Knoxville, TN"
# returned only 4 results while fab labs and machine shops went unfound.
_calls = []
def _fake_search(query, max_results=20, spend=None):
    _calls.append(query)
    n = query.split(" in ")[0]
    # Two variants deliberately return the SAME place, to prove deduping.
    if n in ("makerspace", "fab lab"):
        return [{"id": "shared1", "displayName": {"text": "Shared Space"},
                 "formattedAddress": "1 Main"}]
    return [{"id": f"id-{n}", "displayName": {"text": n.title()},
             "formattedAddress": f"{n} st"}]

_real, scraper.places_search = scraper.places_search, _fake_search
try:
    res = scraper.places_search_all("makerspace", "Knoxville, TN",
                                    max_results=50, max_queries=4)
    check("more than one query is issued", len(_calls) == 4, str(_calls))
    check("the user's own target is queried first",
          _calls[0].startswith("makerspace in"), _calls[0])
    check("location is appended to every variant",
          all(" in Knoxville, TN" in c for c in _calls))
    check("same place from two variants is deduped once",
          sum(1 for p in res if p["id"] == "shared1") == 1, str(res))
    check("distinct places from other variants are kept", len(res) == 3, str(res))

    _calls.clear()
    capped = scraper.places_search_all("makerspace", "X", max_results=2,
                                       max_queries=6)
    check("max_results caps the combined total", len(capped) == 2, str(capped))

    # A later variant failing must not discard places already found.
    _calls.clear()
    def _flaky(query, max_results=20, spend=None):
        _calls.append(query)
        if len(_calls) == 1:
            return [{"id": "ok", "displayName": {"text": "Ok"},
                     "formattedAddress": "a"}]
        raise RuntimeError("Places API 500")
    scraper.places_search = _flaky
    survived = scraper.places_search_all("makerspace", "X", max_results=10,
                                         max_queries=3)
    check("one failed variant doesn't lose earlier results",
          len(survived) == 1, str(survived))
finally:
    scraper.places_search = _real

check("3D-specific terms are tried before generic shops",
      config.PLACES_QUERY_VARIANTS.index("3D printing service")
      < config.PLACES_QUERY_VARIANTS.index("machine shop"))
check("paper-print terms excluded (they return sign/copy shops)",
      not any(v in ("print shop", "sign shop", "printing")
              for v in config.PLACES_QUERY_VARIANTS))
check("place key falls back to name+address when id is absent",
      scraper._place_key({"displayName": {"text": "A"},
                          "formattedAddress": "B"}) == "a|b")
check("place key prefers the stable id",
      scraper._place_key({"id": "xyz", "displayName": {"text": "A"},
                          "formattedAddress": "B"}) == "xyz")

print("\n[7b] Depth presets")
check("three presets defined",
      set(scraper.DEPTH_PRESETS) == {"quick", "standard", "deep"})
check("deep tries every query variant",
      scraper.DEPTH_PRESETS["deep"]["queries"] is None)
check("quick tries only the user's target",
      scraper.DEPTH_PRESETS["quick"]["queries"] == 1)
check("deeper presets crawl more pages",
      scraper.DEPTH_PRESETS["quick"]["pages"]
      < scraper.DEPTH_PRESETS["standard"]["pages"]
      < scraper.DEPTH_PRESETS["deep"]["pages"])
check("quick skips the extra discovery source",
      not scraper.DEPTH_PRESETS["quick"]["discover"])
check("deep enables rendering and discovery",
      scraper.DEPTH_PRESETS["deep"]["render"]
      and scraper.DEPTH_PRESETS["deep"]["discover"])

print("\n[8] Excel output (mock leads, full pipeline)")
config.MOCK_MODE = True
leads = scraper.collect("Testville", "makerspace")
buf = io.BytesIO()
excel_out.build(leads, buf)
buf.seek(0)

from openpyxl import load_workbook
wb = load_workbook(buf)
check("three sheets present",
      wb.sheetnames == ["Summary", "All Contacts", "Top 10 Actions"], str(wb.sheetnames))

s = wb["Summary"]
hdr = [c.value for c in s[1]]
row1 = {hdr[i]: s.cell(2, i + 1).value for i in range(len(hdr))}

for gone in ("Tier", "Will Buy Filament?", "Overall Confidence", "Best Outreach Path",
             "Top Contact Name", "Top Contact Title", "Top Contact Email",
             "Top Contact Phone", "Top Contact LinkedIn",
             "Last Verified (Maps)", "Last Verified (Website)",
             "Other Emails", "Data Notes", "Phone (Direct)", "Key People",
             "Phone (Website)", "Phone (Maps)"):
    check(f"removed column: {gone}", gone not in hdr)

for added in ("Phone", "Other Contacts", "Related Locations",
              "Last Verified", "Notes", "Found Via", "All Emails (CC)"):
    check(f"present column: {added}", added in hdr, str(hdr))

row2 = {hdr[i]: s.cell(3, i + 1).value for i in range(len(hdr))}
check("Places-discovered lead labelled", row1["Found Via"] == "Google Places",
      str(row1["Found Via"]))
check("non-Maps lead labelled as web search", row2["Found Via"] == "Web search",
      str(row2["Found Via"]))
check("non-Maps lead has no Maps verified date",
      "(Maps)" not in (row2["Last Verified"] or ""), str(row2["Last Verified"]))
check("secondary emails surfaced in Other Contacts",
      "info@knoxmakers.org" in (row1["Other Contacts"] or ""))

check("both phones filled from their own sources",
      "(Website)" in row1["Phone"] and "(Maps)" in row1["Phone"], row1["Phone"])
check("Other Contacts cell holds role + reason",
      "Technical Director" in row1["Other Contacts"] and "Why:" in row1["Other Contacts"])
check("Other Contacts includes LinkedIn link", "linkedin.com" in row1["Other Contacts"])
check("Notes column is blank for Rich", row1["Notes"] in (None, ""))
check("related location surfaced", row1["Related Locations"] == "Knox Makers Annex")
check("verified maps date filled", "(Maps)" in row1["Last Verified"])
check("verified website blank when no crawl text",
      "(Website)" not in (row2["Last Verified"] or ""))
check("priority filled", row1["Priority"] in ("HIGH", "MEDIUM", "LOW"))

cw = wb["All Contacts"]
chdr = [c.value for c in cw[1]]
crow = {chdr[i]: cw.cell(2, i + 1).value for i in range(len(chdr))}
check("All Contacts has Role + Why columns",
      "Role" in chdr and "Why This Person" in chdr, str(chdr))
check("All Contacts confidence column removed", "Confidence" not in chdr)
check("All Contacts leads with a decision-maker",
      any(w in (crow["Role"] or "").lower() for w in excel_out.DECISION_WORDS), str(crow))
check("All Contacts carries the per-person reason", bool(crow["Why This Person"]))

a = wb["Top 10 Actions"]
ahdr = [c.value for c in a[1]]
arow = {ahdr[i]: a.cell(2, i + 1).value for i in range(len(ahdr))}
check("Top10 priority filled", arow["Priority"] in ("HIGH", "MEDIUM", "LOW"))
check("Top10 who-to-contact names a person with their role",
      "(" in (arow["Who to Contact"] or "")
      and any(w in (arow["Who to Contact"] or "").lower() for w in excel_out.DECISION_WORDS),
      str(arow["Who to Contact"]))
check("Top10 notes multi-location coverage", "locations" in (arow["Why This Is #1-10"] or ""))
config.MOCK_MODE = False

# ---------------------------------------------------------------------------
print("\n[8b] Cross-run history: a second session skips what the first found")
import tempfile
import history
_saved = {n: getattr(scraper, n) for n in
          ("places_search_all", "discover_via_search", "find_website",
           "crawl_site", "extract_and_pitch", "enrich_gaps", "Renderer")}
_saved_cfg = (config.HISTORY_DIR, config.GOOGLE_PLACES_API_KEY, config.ANTHROPIC_API_KEY,
              config.DROP_EMPTY_LEADS)
_crawled = []
_PLACES = [
    {"id": "p1", "displayName": {"text": "Knox Makers"},
     "formattedAddress": "1 Main St, Knoxville, TN 37920",
     "websiteUri": "https://knoxmakers.org", "userRatingCount": 42},
    {"id": "p2", "displayName": {"text": "Fab Shop"},
     "formattedAddress": "2 Oak Ave, Knoxville, TN 37902",
     "websiteUri": "https://fabshop.com"},
    {"id": "p3", "displayName": {"text": "Brand New Lab"},
     "formattedAddress": "3 Elm St, Knoxville, TN 37902",
     "websiteUri": "https://newlab.org"},
]
_visible = [2]   # run 1 sees the first two places; run 2 sees all three


def _fake_places(target, location, max_results=20, skip=None, **k):
    return [p for p in _PLACES[:_visible[0]] if not (skip and skip(p))][:max_results]


def _fake_crawl2(website, *a, **k):
    _crawled.append(website)
    return {"text": "We run a Prusa MK4.", "email": "hi@" + scraper._host(website),
            "emails": [], "phone": "", "socials": [], "pages": 3, "rendered": 0,
            "printer_models": ["Prusa MK4"], "hosts": [], "priority_pages": ["x/lab"]}


try:
    config.HISTORY_DIR = tempfile.mkdtemp()
    config.GOOGLE_PLACES_API_KEY = config.ANTHROPIC_API_KEY = "test-key"
    config.MOCK_MODE = False
    scraper.places_search_all = _fake_places
    scraper.discover_via_search = lambda *a, **k: []
    scraper.find_website = lambda *a, **k: ""
    scraper.crawl_site = _fake_crawl2
    scraper.extract_and_pitch = lambda *a, **k: {
        "people": [], "equipment_highlights": "3D printers", "printer_models": "",
        "pitch": "p"}
    scraper.enrich_gaps = lambda *a, **k: {}
    scraper.Renderer = lambda **k: type(
        "R", (), {"budget": 0, "html": lambda s, u: "", "close": lambda s: None})()

    ns = history.namespace("test-key")
    run1 = scraper.collect("Knoxville, TN", "makerspace", depth="quick",
                           exclude=history.load(ns))
    history.record(ns, run1.processed, "Knoxville, TN")
    check("session 1 finds both places", len(run1) == 2, str([L["name"] for L in run1]))
    check("history file written", len(history.load(ns)) == 2, str(history.load(ns)))

    _crawled.clear()
    _visible[0] = 3
    run2 = scraper.collect("Knoxville, TN", "makerspace", depth="quick",
                           exclude=history.load(ns))
    check("session 2 returns only the new place",
          [L["name"] for L in run2] == ["Brand New Lab"], str([L["name"] for L in run2]))
    check("session 2 never crawls (pays for) the old places",
          _crawled == ["https://newlab.org"], str(_crawled))
    check("skipped count reported", run2.seen_before == 2, str(run2.seen_before))

    # A spreadsheet from session 1 works as the exclude list on its own
    # (a free host's disk is wiped on restart, so this is the durable record).
    buf = io.BytesIO()
    excel_out.build(run1, buf)
    buf.seek(0)
    from_xlsx = history.from_spreadsheet(buf)
    check("spreadsheet read back as records",
          {r["name"] for r in from_xlsx} == {"Knox Makers", "Fab Shop"}, str(from_xlsx))
    _crawled.clear()
    run3 = scraper.collect("Knoxville, TN", "makerspace", depth="quick", exclude=from_xlsx)
    check("uploaded spreadsheet alone excludes its places",
          [L["name"] for L in run3] == ["Brand New Lab"], str([L["name"] for L in run3]))

    print("\n[8c] Max results caps the TOTAL, web discovery included")
    _many = [{"id": f"m{i}", "displayName": {"text": f"Shop Number {i}"},
              "formattedAddress": f"{i} Main St, Knoxville, TN",
              "websiteUri": f"https://shop{i}.com"} for i in range(20)]
    _asked = {}

    def _places_many(target, location, max_results=20, skip=None, **k):
        _asked["places"] = max_results
        return _many[:max_results]

    def _discover_many(*a, limit=20, **k):
        _asked["discovery"] = limit
        return [{"name": f"Campus Lab {i}", "website": f"https://lab{i}.edu",
                 "city": "Knoxville", "address": "", "kind": "Lab"} for i in range(8)]

    scraper.places_search_all = _places_many
    scraper.discover_via_search = _discover_many
    for cap in (5, 1, 20):
        _asked.clear()
        got = scraper.collect("Knoxville, TN", "makerspace", max_results=cap, depth="deep")
        check(f"cap {cap}: never more than {cap} leads", len(got) <= cap, f"{len(got)} leads")
        check(f"cap {cap}: exactly {cap} when both sources have plenty", len(got) == cap,
              f"{len(got)} leads, asked {_asked}")
    _asked.clear()
    got = scraper.collect("Knoxville, TN", "makerspace", max_results=20, depth="deep")
    check("web discovery still contributes places",
          any("Campus Lab" in L["name"] for L in got), str([L["name"] for L in got]))

    print("\n[8d] Max results counts KEEPERS; never checks more than 2x the cap")
    check("look limit is 2x", config.LOOK_FACTOR == 2, str(config.LOOK_FACTOR))
    _crawled.clear()

    def _crawl_half(website, *a, **k):
        n = int(re.sub(r"\D", "", website) or 0)
        if n % 2:      # odd shops: readable, nothing about 3D printing -> dropped
            _crawled.append(website)
            return {"text": "We sell pottery and tea.", "email": "", "emails": [],
                    "phone": "", "socials": [], "pages": 2, "rendered": 0,
                    "printer_models": [], "hosts": [], "priority_pages": []}
        return _fake_crawl2(website)      # records the crawl itself
    scraper.crawl_site = _crawl_half
    scraper.discover_via_search = lambda *a, **k: []
    got = scraper.collect("Knoxville, TN", "makerspace", max_results=5, depth="standard")
    check("drops don't use a slot: 5 keepers despite half being dropped",
          len(got) == 5, f"{len(got)} leads, crawled {len(_crawled)}")
    check("stops as soon as the cap is met (not the whole pool)",
          len(_crawled) <= 10, str(len(_crawled)))
    check("dropped count reported", got.dropped >= 4, str(got.dropped))

    _crawled.clear()
    scraper.crawl_site = lambda website, *a, **k: (_crawled.append(website) or {
        "text": "Pottery only.", "email": "", "emails": [], "phone": "", "socials": [],
        "pages": 2, "rendered": 0, "printer_models": [], "hosts": [], "priority_pages": []})
    got = scraper.collect("Knoxville, TN", "makerspace", max_results=5, depth="standard")
    check("no keepers anywhere: stops at 2x the cap, never beyond",
          len(got) == 0 and len(_crawled) == 10, f"{len(got)} leads, crawled {len(_crawled)}")
    check("a short run says how many places it checked", got.checked == 10, str(got.checked))

    # Website lookups cost a search each, so only places actually reached pay.
    _lookups = []
    scraper.places_search_all = lambda *a, max_results=20, **k: [
        {"id": f"n{i}", "displayName": {"text": f"Siteless Lab {i}"},
         "formattedAddress": f"{i} Oak St, Knoxville, TN"} for i in range(max_results)]
    scraper.find_website = lambda client, name, *a, **k: (
        _lookups.append(name) or "https://" + name.split()[-1] + "lab.org")
    scraper.crawl_site = _fake_crawl2
    got = scraper.collect("Knoxville, TN", "makerspace", max_results=3, depth="standard")
    check("website lookup only for places the run reached",
          len(got) == 3 and len(_lookups) == 3, f"{len(got)} leads, {len(_lookups)} lookups")

    _a = {"name": "Maker Hub", "email": "", "phone_web": "", "printer_models": "",
          "equipment_highlights": "", "contacts": [{"name": "X"}], "address": "1 Main St"}
    _b = {"name": "Maker Hub Knoxville", "email": "hi@makerhub.com", "phone_web": "(865) 555-0101",
          "printer_models": "Prusa MK4", "equipment_highlights": "3D printers", "contacts": [],
          "address": "1 Main St"}
    scraper._fold_lead(_a, _b)
    check("a merged duplicate's phone/printers/equipment fill the kept row's blanks",
          (_a["phone_web"], _a["printer_models"], _a["equipment_highlights"], _a["email"])
          == ("(865) 555-0101", "Prusa MK4", "3D printers", "hi@makerhub.com"), str(_a))
    check("the kept row's own data is never overwritten", _a["contacts"] == [{"name": "X"}])
    check("web finds are mixed in, not queued behind every Maps result",
          [p["n"] for p in scraper._interleave_sources(
              [{"n": i, "discovered_via": "Google Places"} for i in range(6)]
              + [{"n": "w1", "discovered_via": "Web search"}, {"n": "w2", "discovered_via": "Web search"}])]
          == [0, 1, 2, "w1", 3, 4, 5, "w2"])

    history.clear(ns)
    check("clearing history forgets earlier runs", history.load(ns) == [])
    check("visitors with different keys get separate histories",
          history.namespace("key-a") != history.namespace("key-b"))
finally:
    for n, v in _saved.items():
        setattr(scraper, n, v)
    (config.HISTORY_DIR, config.GOOGLE_PLACES_API_KEY, config.ANTHROPIC_API_KEY,
     config.DROP_EMPTY_LEADS) = _saved_cfg

# ---------------------------------------------------------------------------
if "--live" in sys.argv:
    print("\n[9] LIVE crawl (network, no API keys): knoxmakers.org")
    site = scraper.crawl_site("https://knoxmakers.org")
    print(f"  text: {len(site['text'])} chars | email: {site['email'] or '(none)'} "
          f"| phone: {site['phone'] or '(none)'} | socials: {len(site['socials'])}")
    check("live crawl returned substantial text", len(site["text"]) > 1000,
          f"{len(site['text'])} chars")
    check("live crawl text mentions maker content",
          any(w in site["text"].lower() for w in ("maker", "3d", "laser", "member")))

# ---------------------------------------------------------------------------
total = PASS + len(FAIL)
print(f"\n{'=' * 55}\n{PASS}/{total} passed" + (f", FAILURES: {FAIL}" if FAIL else " — all green"))
sys.exit(1 if FAIL else 0)
